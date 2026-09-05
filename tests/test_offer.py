"""Stufe 6: Angebotsentwurf - ein Arbeitsdokument, kein Angebot."""

from __future__ import annotations

from datetime import UTC, datetime
from pathlib import Path

from tender_ai.models.calculation import PositionCost, TenderCalculation
from tender_ai.models.decision import PLACEHOLDER, DecisionKind, UserDecision
from tender_ai.models.tender import Tender
from tender_ai.offer import (
    DRAFT_BANNER,
    build_draft,
    draft_filename,
    render_markdown,
    write_markdown,
    write_xlsx,
)


def _tender() -> Tender:
    return Tender(
        id="fixture:1",
        source="fixture",
        source_id="1",
        title="Lieferung von Monitoren",
        contracting_authority="Stadt Musterhausen",
        national_id="VG-2026-1",
        submission_deadline=datetime(2036, 11, 15, 12, 0, tzinfo=UTC),
    )


def _calculation() -> TenderCalculation:
    calculation = TenderCalculation(tender_id="fixture:1", currency="EUR")
    calculation.positions = [
        PositionCost(
            position="1.10",
            title="Monitor 27 Zoll",
            quantity=100.0,
            unit="STK",
            supplier="Muster Distribution",
            match_confidence=90,
            unit_purchase_price=180.0,
            cost_total=18000.0,
            sale_total=22500.0,
        ),
        PositionCost(
            position="1.20",
            title="Sonderanfertigung",
            quantity=5.0,
            unit="STK",
            warnings=["Kein belastbar zugeordnetes Angebot"],
        ),
    ]
    calculation.review_notes = ["Diese Rechnung ist eine Entscheidungsvorlage, kein Angebot."]
    return calculation


def _decision() -> UserDecision:
    return UserDecision(
        tender_id="fixture:1",
        kind=DecisionKind.APPROVED,
        decided_by="j.fichtner",
        decided_at=datetime(2026, 9, 5, 10, 0, tzinfo=UTC),
    )


def test_draft_keeps_unpriced_positions_as_manual_rows():
    """Eine still weggelassene Position waere im Angebot ein fehlender Posten."""
    draft = build_draft(_tender(), _calculation(), _decision())

    assert len(draft.positions) == 2
    assert draft.manual_position_count == 1
    manual = draft.positions[1]
    assert manual.needs_manual_entry
    assert manual.total_price is None
    assert manual.source_note  # der Grund steht dabei


def test_draft_total_covers_only_priced_positions():
    draft = build_draft(_tender(), _calculation(), _decision())
    assert draft.net_total == 22500.0
    assert any("ohne Preis" in note for note in draft.review_notes)


def test_draft_carries_the_approval():
    draft = build_draft(_tender(), _calculation(), _decision())
    assert draft.approved_by == "j.fichtner"
    assert draft.approved_at is not None


def test_draft_never_invents_bidder_data():
    """Was das Tool nicht wissen kann, bleibt Platzhalter."""
    draft = build_draft(_tender(), _calculation(), _decision())
    assert draft.placeholders
    assert any("Unterschrift" in item for item in draft.placeholders)
    assert any("Bieter" in item for item in draft.placeholders)


def test_markdown_marks_the_document_as_a_draft_everywhere():
    text = render_markdown(build_draft(_tender(), _calculation(), _decision()))

    assert text.startswith(f"# {DRAFT_BANNER}")
    assert text.rstrip().endswith(f"**{DRAFT_BANNER}**")
    assert "kein Angebot" in text
    assert PLACEHOLDER in text
    # Kein Preis wird erfunden: die unbepreiste Zeile traegt den Platzhalter.
    assert "Sonderanfertigung" in text
    assert text.count(PLACEHOLDER) >= len(build_draft(_tender(), _calculation()).placeholders)


def test_markdown_shows_where_each_price_comes_from():
    text = render_markdown(build_draft(_tender(), _calculation(), _decision()))
    assert "Muster Distribution" in text
    assert "Zuordnungsguete 90" in text


def test_files_are_written_with_a_draft_filename(tmp_path: Path):
    draft = build_draft(_tender(), _calculation(), _decision())
    markdown = write_markdown(
        draft, tmp_path / draft_filename("fixture:1", ".md", draft.generated_at)
    )
    workbook = write_xlsx(
        draft, tmp_path / draft_filename("fixture:1", ".xlsx", draft.generated_at)
    )

    assert markdown.name.startswith("ENTWURF-")
    assert workbook.name.startswith("ENTWURF-")
    assert markdown.exists() and workbook.exists()


def test_xlsx_states_the_draft_status_in_the_first_row(tmp_path: Path):
    """Ein Preisblatt wandert weiter als das Anschreiben."""
    from openpyxl import load_workbook

    draft = build_draft(_tender(), _calculation(), _decision())
    path = write_xlsx(draft, tmp_path / "entwurf.xlsx")

    sheet = load_workbook(path).active
    assert sheet.title == "ENTWURF"
    assert sheet["A1"].value == DRAFT_BANNER
    values = [row[0] for row in sheet.iter_rows(values_only=True)]
    assert DRAFT_BANNER in values[1:]  # auch am Ende


def test_filename_is_filesystem_safe():
    name = draft_filename("ted:00123456-2026", ".md", datetime(2026, 9, 5, 10, 0, tzinfo=UTC))
    assert ":" not in name
    assert name == "ENTWURF-ted-00123456-2026-20260905-1000.md"


def test_machine_output_marks_the_draft_as_unsubmitted():
    payload = build_draft(_tender(), _calculation(), _decision()).as_dict()
    assert payload["is_draft"] is True
    assert payload["is_submitted"] is False
    assert payload["requires_manual_submission"] is True
