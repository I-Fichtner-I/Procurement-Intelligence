"""Angebotsentwurf erzeugen (Stufe 6).

Ein Entwurf ist ein Arbeitsdokument. Er wird nirgends eingereicht, an niemanden
versendet und traegt keine Unterschrift. Damit er nicht versehentlich als
fertiges Angebot behandelt wird, macht das Dokument seinen Zustand an jeder
Stelle sichtbar: Kopfzeile, Fusszeile und auffaellige Platzhalter fuer alles,
was das Tool nicht wissen kann.

Was das Tool nicht wissen kann, erfindet es auch nicht: Bieterangaben,
Zahlungsbedingungen, Nachlaesse und Unterschrift bleiben Platzhalter.
"""

from __future__ import annotations

from datetime import datetime
from pathlib import Path

from ..models.calculation import TenderCalculation
from ..models.common import display
from ..models.decision import PLACEHOLDER, DraftPosition, OfferDraft, UserDecision
from ..models.tender import Tender

#: Angaben, die ein Mensch ergaenzen muss - das Tool kennt sie nicht.
REQUIRED_PLACEHOLDERS = (
    "Bieter: Firma, Anschrift, Ansprechpartner",
    "Umsatzsteuer-Identifikationsnummer",
    "Zahlungsbedingungen und Skonto",
    "Bindefrist des Angebots",
    "Nachlaesse und Rabatte",
    "Ort, Datum, rechtsverbindliche Unterschrift",
)

#: Steht in Kopf- und Fusszeile jedes erzeugten Dokuments.
DRAFT_BANNER = "ENTWURF - NICHT ZUR ABGABE"


def build_draft(
    tender: Tender,
    calculation: TenderCalculation,
    decision: UserDecision | None = None,
) -> OfferDraft:
    """Aus Ausschreibung und Kalkulation einen Entwurf bauen.

    Positionen ohne belastbaren Preis verschwinden nicht - sie stehen als
    Zeilen mit Handarbeitskennzeichen im Entwurf. Eine stillschweigend
    weggelassene Position waere im abgegebenen Angebot ein fehlender Posten.
    """
    draft = OfferDraft(
        tender_id=calculation.tender_id,
        tender_title=tender.title,
        contracting_authority=tender.contracting_authority,
        national_id=tender.national_id,
        submission_deadline=tender.submission_deadline,
        currency=calculation.currency,
        placeholders=list(REQUIRED_PLACEHOLDERS),
        approved_by=decision.decided_by if decision else None,
        approved_at=decision.decided_at if decision else None,
    )

    net_total = 0.0
    has_price = False
    for position in calculation.positions:
        unit_price = None
        if position.sale_total is not None and position.quantity:
            unit_price = position.sale_total / position.quantity
        draft_position = DraftPosition(
            position=position.position,
            title=position.title,
            quantity=position.quantity,
            unit=position.unit,
            unit_price=unit_price,
            total_price=position.sale_total,
            needs_manual_entry=position.sale_total is None,
        )
        if position.sale_total is None:
            draft_position.source_note = (
                position.warnings[0] if position.warnings else "Kein belastbarer Preis"
            )
        else:
            has_price = True
            net_total += position.sale_total
            draft_position.source_note = (
                f"Einkauf {position.supplier or display(None)}, "
                f"Zuordnungsguete {position.match_confidence}"
            )
        draft.positions.append(draft_position)

    draft.net_total = net_total if has_price else None
    draft.review_notes = list(calculation.review_notes)
    if draft.manual_position_count:
        draft.review_notes.insert(
            0,
            f"{draft.manual_position_count} Position(en) ohne Preis - vor der Abgabe "
            f"von Hand ergaenzen. Die Summe unten ist ohne sie gerechnet.",
        )
    return draft


def _money(amount: float | None, currency: str | None) -> str:
    if amount is None:
        return PLACEHOLDER
    text = f"{amount:,.2f}".replace(",", "@").replace(".", ",").replace("@", ".")
    return f"{text} {currency}" if currency else text


def render_markdown(draft: OfferDraft) -> str:
    """Entwurf als Markdown - lesbar, versionierbar, leicht zu korrigieren."""
    lines: list[str] = [
        f"# {DRAFT_BANNER}",
        "",
        "> Dieses Dokument ist ein maschinell erzeugter Entwurf. Es wurde nicht",
        "> eingereicht und ist kein Angebot. Vor einer Abgabe sind alle mit",
        f"> `{PLACEHOLDER}` markierten Angaben zu ergaenzen und saemtliche",
        "> Positionen, Mengen und Preise von Hand zu pruefen.",
        "",
        "## Ausschreibung",
        "",
        f"- **Titel:** {display(draft.tender_title)}",
        f"- **Vergabestelle:** {display(draft.contracting_authority)}",
        f"- **Amtliche Nummer:** {display(draft.national_id)}",
        f"- **Angebotsfrist:** {display(draft.submission_deadline)}",
        f"- **Interne Kennung:** {draft.tender_id}",
        "",
        "## Bieter",
        "",
    ]
    lines.extend(f"- **{item}:** {PLACEHOLDER}" for item in draft.placeholders)
    lines.extend(
        [
            "",
            "## Positionen",
            "",
            "| Pos. | Bezeichnung | Menge | Einheit | Einzelpreis | Gesamtpreis | Herkunft |",
            "|------|-------------|-------|---------|-------------|-------------|----------|",
        ]
    )
    for position in draft.positions:
        marker = " **(von Hand)**" if position.needs_manual_entry else ""
        lines.append(
            f"| {display(position.position)} | {position.title}{marker} "
            f"| {display(position.quantity)} | {display(position.unit)} "
            f"| {_money(position.unit_price, draft.currency)} "
            f"| {_money(position.total_price, draft.currency)} "
            f"| {display(position.source_note)} |"
        )

    lines.extend(
        [
            "",
            f"**Nettosumme der bepreisten Positionen:** {_money(draft.net_total, draft.currency)}",
            "",
            f"**Umsatzsteuer:** {PLACEHOLDER}",
            "",
            f"**Bruttosumme:** {PLACEHOLDER}",
            "",
            "## Vor der Abgabe pruefen",
            "",
        ]
    )
    lines.extend(f"- {note}" for note in draft.review_notes)
    lines.extend(
        [
            "",
            "---",
            "",
            f"Erzeugt am {draft.generated_at:%Y-%m-%d %H:%M} UTC"
            + (
                f" nach Freigabe durch {draft.approved_by}"
                f" am {draft.approved_at:%Y-%m-%d %H:%M} UTC"
                if draft.approved_by and draft.approved_at
                else ""
            ),
            "",
            f"**{DRAFT_BANNER}**",
            "",
        ]
    )
    return "\n".join(lines)


def write_markdown(draft: OfferDraft, destination: Path) -> Path:
    destination.parent.mkdir(parents=True, exist_ok=True)
    destination.write_text(render_markdown(draft), encoding="utf-8")
    return destination


def write_xlsx(draft: OfferDraft, destination: Path) -> Path:
    """Entwurf als Tabelle - viele Vergabestellen wollen ein Preisblatt.

    Auch hier steht der Entwurfsstatus in der ersten Zeile und im Blattnamen:
    ein Preisblatt wandert erfahrungsgemaess weiter als das Anschreiben.
    """
    from openpyxl import Workbook
    from openpyxl.styles import Font

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "ENTWURF"

    sheet.append([DRAFT_BANNER])
    sheet["A1"].font = Font(bold=True, color="B00020")
    sheet.append(["Kein Angebot. Vor der Abgabe von Hand pruefen und ergaenzen."])
    sheet.append([])
    sheet.append(["Ausschreibung", display(draft.tender_title)])
    sheet.append(["Vergabestelle", display(draft.contracting_authority)])
    sheet.append(["Amtliche Nummer", display(draft.national_id)])
    sheet.append(["Angebotsfrist", display(draft.submission_deadline)])
    sheet.append([])

    header = ["Pos.", "Bezeichnung", "Menge", "Einheit", "Einzelpreis", "Gesamtpreis", "Herkunft"]
    sheet.append(header)
    for cell in sheet[sheet.max_row]:
        cell.font = Font(bold=True)

    for position in draft.positions:
        sheet.append(
            [
                display(position.position),
                position.title + (" (von Hand)" if position.needs_manual_entry else ""),
                position.quantity,
                display(position.unit),
                position.unit_price,
                position.total_price,
                display(position.source_note),
            ]
        )

    sheet.append([])
    sheet.append(["Nettosumme", None, None, None, None, draft.net_total])
    sheet.append(["Umsatzsteuer", None, None, None, None, PLACEHOLDER])
    sheet.append(["Bruttosumme", None, None, None, None, PLACEHOLDER])
    sheet.append([])
    for item in draft.placeholders:
        sheet.append([item, PLACEHOLDER])
    sheet.append([])
    sheet.append([DRAFT_BANNER])
    sheet[f"A{sheet.max_row}"].font = Font(bold=True, color="B00020")

    for column, width in zip("ABCDEFG", (10, 52, 10, 10, 14, 14, 44), strict=False):
        sheet.column_dimensions[column].width = width

    destination.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(destination)
    return destination


#: Jeder Entwurf traegt es im Namen - eine Datei, die aussieht wie ein Angebot,
#: soll sich schon im Dateimanager als Entwurf zu erkennen geben.
DRAFT_PREFIX = "ENTWURF-"


def safe_stem(tender_id: str) -> str:
    """Tender-ID als Dateinamensteil: nur Buchstaben, Ziffern, Strich, Unterstrich."""
    return "".join(char if char.isalnum() or char in "-_" else "-" for char in tender_id)


def draft_name_prefix(tender_id: str) -> str:
    """Anfang aller Entwurfsdateien dieser Ausschreibung - zum Wiederfinden."""
    return f"{DRAFT_PREFIX}{safe_stem(tender_id)}-"


def draft_filename(tender_id: str, suffix: str, generated_at: datetime) -> str:
    """Dateiname mit Entwurfskennzeichen und Zeitstempel."""
    return f"{draft_name_prefix(tender_id)}{generated_at:%Y%m%d-%H%M}{suffix}"
